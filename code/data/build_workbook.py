"""Build data/dataset_overview.xlsx: Analysis (charts) + readable data sheets for train and val.

Sheets: Analysis, Calls, Facts, Transcripts, Noise, Summaries, Validator.
Everything is derived from the authored YAML (the single source of truth); never edit the workbook by hand.

Run: python3 code/data/build_workbook.py
"""
from __future__ import annotations

import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from common.loader import ROOT, load_split  # noqa: E402
from common.render import render  # noqa: E402
from data.validate_gold import global_checks, validate_call  # noqa: E402

OUT = ROOT / "data" / "dataset_overview.xlsx"
SPLITS = ["train", "val"]
LABEL = {"train": "Training", "val": "Validation"}
BLUE, ORANGE = "4472C4", "ED7D31"
SPLIT_FILL = {"train": PatternFill("solid", fgColor="DDEBF7"), "val": PatternFill("solid", fgColor="FCE4D6")}
HEAD = PatternFill("solid", fgColor="1F3864")
NOISE_FILL = PatternFill("solid", fgColor="FFF2CC")
LABEL_ERR_FILL = PatternFill("solid", fgColor="F8CBAD")
WRAP = Alignment(wrap_text=True, vertical="top")

CATEGORY_ORDER = ["routine", "ambiguous", "asr_error", "medication", "supply", "high_risk", "not_applicable"]
CATEGORY_TARGET = {"train": [80, 60, 60, 90, 60, 100, 50], "val": [16, 12, 12, 18, 12, 20, 10]}
RULES = {
    "G1": "Schema: valid category, length bucket, ids, flags",
    "G2": "Every summary quote exists verbatim in the cited turn, said by the cited speaker",
    "G3": "Numbers in the summary were spoken in the evidence turn",
    "G4": "Every fact, flag and education item is tagged on a transcript turn",
    "G5": "Identity (name, DOB, phone, relationship) matches what was spoken; frequency strings match",
    "G6": "Summary facts match the fact record (medication, symptom, action, education)",
    "G7": "Medication fields in the summary match the record",
    "G8": "Planned vs completed action wording is consistent",
    "G9": "Certainty (stated / unclear) is consistent with the transcript",
    "G10": "Every record item is covered by a summary bullet and vice versa",
    "G11": "Drug names in untagged turns are declared as distractors",
    "G12": "Length bucket (chars and line count) matches the v2 regime",
}


def has_tag(c, *tags):
    return bool(set(c["setup"].get("secondary_tags") or []) & set(tags))


def noise_types(c):
    return {e["type"] for e in c["noise_events"]}


def any_item(c, key, pred):
    return any(pred(i) for i in c[key])


# (checklist item, predicate over a loaded call) — the user's coverage list
CHECKLIST = [
    ("Medication name + dose", lambda c: any_item(c, "assessment", lambda i: i.get("item_type") == "medication" and i.get("dose"))),
    ("Ambiguity (unclear certainty)", lambda c: any_item(c, "assessment", lambda i: i.get("certainty") == "unclear")),
    ("Explicit self-correction", lambda c: has_tag(c, "dob_self_corrected", "self_correction", "explicit_self_correction") or "correction" in str(c["setup"].get("subtype"))),
    ("Unresolved conflict", lambda c: "unresolved" in str(c["setup"].get("subtype")) or has_tag(c, "identity_unresolved", "contradiction")),
    ("PRN + scheduled medication", lambda c: has_tag(c, "prn_and_scheduled") or "prn_and_scheduled" in str(c["setup"].get("subtype"))),
    ("Dose / concentration errors", lambda c: has_tag(c, "liquid_concentration", "wrong_dose", "extra_dose") or "dose" in str(c["setup"].get("subtype")) or any_item(c, "assessment", lambda i: i.get("medication_status") == "extra_dose")),
    ("DOB / name / number errors", lambda c: bool(noise_types(c) & {"number_garble", "name_drift"}) or has_tag(c, "name_drift", "dob_self_corrected", "number_garble", "dob_not_stated")),
    ("Hedging language", lambda c: has_tag(c, "hedging")),
    ("Planned actions", lambda c: any_item(c, "response", lambda i: i.get("action_status") == "planned")),
    ("Completed actions", lambda c: any_item(c, "response", lambda i: i.get("action_status") == "completed")),
    ("Pertinent negatives", lambda c: any_item(c, "assessment", lambda i: i.get("item_type") == "pertinent_negative")),
    ("High-risk situation", lambda c: c["setup"].get("scenario_category") == "high_risk"),
    ("Risk flags raised", lambda c: bool(c["risk_flags"])),
    ("Not applicable (non-clinical)", lambda c: c["setup"].get("scenario_category") == "not_applicable"),
    ("Multi-issue call", lambda c: has_tag(c, "multi_issue")),
    ("Long / messy call", lambda c: c["setup"].get("length_bucket") == "long"),
    ("ASR corruption", lambda c: bool(c["noise_events"]) or c["setup"].get("scenario_category") == "asr_error"),
    ("Drug-name garble", lambda c: "drug_garble" in noise_types(c)),
    ("Late / split turns", lambda c: has_tag(c, "late_answers", "split_turns")),
    ("Irrelevant chatter", lambda c: has_tag(c, "offtopic_chatter") or bool(c["setup"].get("offtopic_chatter"))),
    ("Speaker attribution errors", lambda c: "speaker_label" in noise_types(c) or has_tag(c, "speaker_label_errors", "crosstalk")),
    ("Trap / negation phrasing", lambda c: has_tag(c, "trap_phrasing", "negated_risk_phrase", "negated_risk")),
    ("Look-alike (tall-man) drugs", lambda c: has_tag(c, "tall_man_drug", "tall_man")),
    ("Name spelled out", lambda c: has_tag(c, "name_spelled")),
    ("Phone missing / read from file", lambda c: has_tag(c, "phone_not_given", "phone_from_file")),
    ("Figurative risk language", lambda c: has_tag(c, "figurative_risk_language")),
]


# ---------------------------------------------------------------- helpers
def style_header(ws, row=1, color=HEAD):
    for cell in ws[row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = color
            cell.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, spec):
    for i, w in enumerate(spec, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def clean(v):
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x) for x in v)
    return v


def headline(item: dict, kind: str) -> tuple[str, str]:
    """(short readable headline, remaining details) for an assessment/response/education row."""
    skip = {"item_id", "item_type", "certainty", "reporter", "actor", "speaker"}
    used = set()

    def g(*keys):
        parts = []
        for k in keys:
            v = item.get(k)
            if v not in (None, "", [], False):
                parts.append(str(v) if v is not True else k)
                used.add(k)
        return parts

    if kind == "assessment":
        t = item.get("item_type")
        if t == "symptom":
            used.update(["symptom_name", "symptom_present"])
            head = f"{item.get('symptom_name')} ({'present' if item.get('symptom_present') else 'absent'})"
        elif t == "pertinent_negative":
            used.update(["symptom_name", "symptom_present"])
            head = f"No {item.get('symptom_name')}"
        elif t == "medication":
            name = item.get("medication_name_spoken") or item.get("medication_name_true")
            used.update(["medication_name_spoken", "medication_name_true"])
            head = " ".join([str(name)] + g("strength", "dose", "unit", "route", "frequency") + (["PRN"] if item.get("prn") else []))
            if item.get("prn"):
                used.add("prn")
            st = item.get("medication_status")
            if st:
                used.add("medication_status")
                head += f" [{st}]"
        elif t == "supply":
            head = " ".join(g("supply_item", "supply_remaining", "supply_unit"))
        elif t == "vital":
            head = " ".join(g("vital_name", "vital_value"))
        else:
            head = " ".join(g("history_or_context_text", "caller_concern_text"))
    elif kind == "response":
        head = " ".join(g("action_type", "action_status", "recipient_or_target", "timeframe"))
    else:
        head = " ".join(g("education_type")) + (": " + str(item.get("instruction_content", "")) if item.get("instruction_content") else "")
        used.add("instruction_content")
    rest = "; ".join(f"{k}: {clean(v)}" for k, v in item.items()
                     if k not in skip and k not in used and v not in (None, "", [], False))
    return head.strip(), rest


def evidence(call):
    """item_id -> (speaker, quote, turns) from the gold alignment."""
    out = {}
    tgt = call["target"]
    for label, ids in call["alignment"].items():
        if label == "chief_complaint":
            node = tgt.get("chief_complaint")
        else:
            sec, idx = label[:-1].split("[")
            node = (tgt.get(sec) or [])[int(idx)]
        if not node:
            continue
        for i in ids:
            out[i] = (node.get("speaker"), node.get("quote"), clean(node.get("turns")))
    return out


# ---------------------------------------------------------------- data sheets
def sheet_calls(wb, by_split):
    ws = wb.create_sheet("Calls")
    ws.append(["Call", "Split", "Category", "Subtype", "Agency", "Agency type", "Caller relationship", "Patient age",
               "Length", "Turns", "Characters", "ASR noise", "Noise events", "Expected risk flags", "Assessment items",
               "Response items", "Education items", "Scenario tags", "Patient (true)", "Caller", "Reason for call", "Scenario blueprint"])
    for split in SPLITS:
        for c in by_split[split]:
            s, cc = c["setup"], c["chief_complaint"]
            ws.append([c["call_id"], LABEL[split], s.get("scenario_category"), s.get("subtype"), s.get("agency_name"), s.get("agency_type"),
                       s.get("relationship_group"), s.get("patient_age_group"), s.get("length_bucket"), len(c["turns"]), len(c["transcript"]),
                       s.get("asr_noise_level"), len(c["noise_events"]), clean([f["risk_category"] for f in c["risk_flags"]]),
                       len(c["assessment"]), len(c["response"]), len(c["education"]), clean(s.get("secondary_tags")),
                       cc.get("patient_name_true"), cc.get("caller_name_true"), cc.get("reason_for_call"), c["blueprint"]])
            ws.cell(ws.max_row, 2).fill = SPLIT_FILL[split]
    style_header(ws)
    widths(ws, [9, 11, 15, 26, 28, 14, 17, 11, 9, 7, 10, 10, 9, 28, 10, 10, 10, 40, 22, 22, 50, 80])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def sheet_facts(wb, by_split):
    ws = wb.create_sheet("Facts")
    ws.append(["Call", "Split", "Section", "Item", "Type", "Fact", "Details", "Certainty", "Said by", "Evidence quote", "Turn(s)"])
    for split in SPLITS:
        for c in by_split[split]:
            ev = evidence(c)
            cc = c["chief_complaint"]
            rows = []
            if c["setup"].get("scenario_category") == "not_applicable":
                rows.append(["Identity", "-", "not applicable", c["setup"].get("na_reason") or "", "", "", "", "", "", ""])
            else:
                for label, key in (("Patient name", "patient_name"), ("Date of birth", "patient_dob"), ("Caller name", "caller_name"), ("Callback phone", "callback_phone")):
                    heard = clean(cc.get(f"{key}_heard_as"))
                    spoken = cc.get(f"{key}_spoken")
                    rows.append(["Identity", key, label, cc.get(f"{key}_true") or "(not stated)",
                                 f"spoken: {spoken}" + (f"; heard as: {heard}" if heard else ""), cc.get(f"{key}_certainty"), "", "", ""])
                rows.append(["Identity", "relationship", "Relationship", cc.get("relationship_to_patient") or "", "", cc.get("relationship_certainty"), "", "", ""])
            rows.append(["Chief complaint", "cc", cc.get("reason_category") or "", cc.get("reason_for_call"), cc.get("relevant_context") or "", "stated", *ev.get("cc", ("", "", ""))])
            for key, sec in (("assessment", "Assessment"), ("response", "Response"), ("education", "Education")):
                for it in c[key]:
                    head, rest = headline(it, key)
                    typ = it.get("item_type") or it.get("action_type") or it.get("education_type")
                    rows.append([sec, it["item_id"], typ, head, rest, it.get("certainty"), *ev.get(it["item_id"], ("", "", ""))])
            for f in c["risk_flags"]:
                rows.append(["Risk flag", f["flag_id"], f["risk_category"], f.get("trigger_description"), f"required: {f.get('flag_required')}", "", f.get("trigger_speaker"), "", ""])
            for r in rows:
                section, item, typ, fact, details, cert, who, quote, turns = (r + [""] * 9)[:9]
                ws.append([c["call_id"], LABEL[split], section, item, typ, fact, details, cert, who, quote, turns])
                ws.cell(ws.max_row, 2).fill = SPLIT_FILL[split]
                if cert == "unclear":
                    ws.cell(ws.max_row, 8).fill = NOISE_FILL
    style_header(ws)
    widths(ws, [9, 11, 15, 8, 24, 55, 45, 10, 9, 60, 9])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def sheet_transcripts(wb, by_split):
    ws = wb.create_sheet("Transcripts")
    ws.append(["Call", "Split", "Turn", "Speaker", "Text (as the ASR heard it)", "Fact tags", "Noise logged on this turn"])
    for split in SPLITS:
        for c in by_split[split]:
            notes = defaultdict(list)
            for e in c["noise_events"]:
                notes[e["turn"]].append(f"{e['type']}: said \"{e.get('true_form')}\" -> heard \"{e.get('heard_form')}\"")
            for t in c["turns"]:
                note = " | ".join(notes.get(t["id"], []))
                ws.append([c["call_id"], LABEL[split], t["id"], t["speaker"], t["text"], clean(t["fact_ids"]), note])
                if note or t["label_error"]:
                    for col in range(3, 8):
                        ws.cell(ws.max_row, col).fill = LABEL_ERR_FILL if t["label_error"] else NOISE_FILL
                ws.cell(ws.max_row, 2).fill = SPLIT_FILL[split]
    style_header(ws)
    widths(ws, [9, 11, 6, 9, 110, 22, 60])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def sheet_noise(wb, by_split):
    ws = wb.create_sheet("Noise")
    ws.append(["Call", "Split", "Category", "ASR noise level", "Turn", "Noise type", "What was said", "What the ASR heard", "Linked fact", "Transcript line"])
    for split in SPLITS:
        for c in by_split[split]:
            line = {t["id"]: t for t in c["turns"]}
            for e in c["noise_events"]:
                t = line.get(e["turn"], {})
                ws.append([c["call_id"], LABEL[split], c["setup"].get("scenario_category"), c["setup"].get("asr_noise_level"), e["turn"], e["type"],
                           e.get("true_form"), e.get("heard_form"), e.get("fact_id"), f"{t.get('speaker', '')} -> {t.get('text', '')}"])
                ws.cell(ws.max_row, 2).fill = SPLIT_FILL[split]
    style_header(ws)
    widths(ws, [9, 11, 15, 10, 6, 15, 40, 40, 10, 100])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def bullets(items):
    out = []
    for it in items or []:
        q = f'\n    {it.get("speaker")}: "{it.get("quote")}"' if it.get("quote") else ""
        out.append(f"• {it.get('text')}{q}")
    return "\n".join(out)


def sheet_summaries(wb, by_split):
    ws = wb.create_sheet("Summaries")
    ws.append(["Call", "Split", "Category", "Patient", "Chief complaint", "Assessment", "Response", "Education",
               "Risk flags (returned outside the summary)", "Rendered summary (reference format)"])
    for split in SPLITS:
        for c in by_split[split]:
            t = c["target"]
            ts = c["setup"].get("call_timestamp_utc")
            if t["not_applicable"]["is_na"]:
                ws.append([c["call_id"], LABEL[split], c["setup"].get("scenario_category"), "-", "NOT APPLICABLE: " + (t["not_applicable"].get("reason") or ""),
                           "", "", "", "", render(t, ts)])
            else:
                ident = t["identity"]
                cc = t["chief_complaint"]
                flags = "\n".join(f"{f.get('risk_category')}: {f.get('trigger_description')}" for f in c["risk_flags"])
                ws.append([c["call_id"], LABEL[split], c["setup"].get("scenario_category"),
                           f"{ident['patient_name']['value']} (DOB {ident['patient_dob']['value'] or 'not stated'})",
                           f"{cc['reason']}\n    {cc['speaker']}: \"{cc['quote']}\"", bullets(t["assessment"]), bullets(t["response"]),
                           bullets(t["education"]), flags, render(t, ts)])
            ws.cell(ws.max_row, 2).fill = SPLIT_FILL[split]
    style_header(ws)
    widths(ws, [9, 11, 15, 28, 55, 70, 60, 60, 45, 90])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def sheet_validator(wb, by_split):
    ws = wb.create_sheet("Validator")
    ws.append(["Rule", "What it checks", "Training violations", "Validation violations"])
    counts = {s: Counter() for s in SPLITS}
    calls_with = {s: 0 for s in SPLITS}
    for s in SPLITS:
        for c in by_split[s]:
            items = validate_call(c).items
            calls_with[s] += bool(items)
            for lvl, code, _ in items:
                counts[s][code] += 1
    for code, text in RULES.items():
        ws.append([code, text, counts["train"][code], counts["val"][code]])
    leaks = global_checks([c for s in SPLITS for c in by_split[s]])
    ws.append([])
    ws.append(["LEAK", "Name / DOB / phone shared between Training and Validation", sum(l.startswith("ERROR") for l in leaks), ""])
    ws.append([])
    ws.append(["", "Calls checked", len(by_split["train"]), len(by_split["val"])])
    ws.append(["", "Calls with any finding", calls_with["train"], calls_with["val"]])
    style_header(ws)
    widths(ws, [8, 90, 20, 22])
    return counts, calls_with, leaks


# ---------------------------------------------------------------- analysis
class Board:
    """Lays out table + chart blocks down the Analysis sheet."""

    def __init__(self, ws):
        self.ws, self.row = ws, 8

    def table(self, title, header, rows, chart="bar", note=None, anchor_col=7, series_cols=None, height=15, width=19, stacked=False):
        ws, r0 = self.ws, self.row
        ws.cell(r0, 1, title).font = Font(bold=True, size=13, color="1F3864")
        if note:
            ws.cell(r0 + 1, 1, note).font = Font(italic=True, color="595959")
        h = r0 + 2
        for j, name in enumerate(header):
            c = ws.cell(h, 1 + j, name)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                ws.cell(h + 1 + i, 1 + j, v)
        last = h + len(rows)
        if chart:
            n = len(header) - 1
            if chart == "pie":
                ch = PieChart()
                ch.add_data(Reference(ws, min_col=2, min_row=h, max_row=last), titles_from_data=True)
                ch.dataLabels = DataLabelList()
                ch.dataLabels.showPercent = True
            else:
                ch = BarChart()
                ch.type = "bar" if len(rows) > 8 else "col"
                if stacked:
                    ch.grouping, ch.overlap = "stacked", 100
                cols = series_cols or range(2, 2 + n)
                for k in cols:
                    ch.add_data(Reference(ws, min_col=k, min_row=h, max_row=last), titles_from_data=True)
                colors = [BLUE, ORANGE, "A5A5A5", "FFC000"]
                for s, col in zip(ch.series, colors):
                    s.graphicalProperties.solidFill = col
                ch.y_axis.majorGridlines = None
                ch.y_axis.delete = False
                ch.x_axis.delete = False
            ch.set_categories(Reference(ws, min_col=1, min_row=h + 1, max_row=last))
            ch.title = title
            ch.height, ch.width = height, width
            ws.add_chart(ch, f"{get_column_letter(anchor_col)}{r0}")
        self.row = max(last + 3, r0 + (int(height * 1.95) if chart else 0) + 2)


def sheet_analysis(wb, by_split, counts, calls_with, leaks):
    ws = wb.create_sheet("Analysis", 0)
    ws.sheet_view.showGridLines = False
    widths(ws, [34, 14, 14, 14, 14, 4] + [11] * 16)
    ws["A1"] = "Clinical Summary Dataset: overview"
    ws["A1"].font = Font(bold=True, size=20, color="1F3864")
    ws["A2"] = "Hand-authored phone-call transcripts (fact record -> transcript with logged ASR noise -> gold summary), every call checked by the gold validator. Training and Validation share no patient, caller, DOB or phone."
    ws["A2"].font = Font(italic=True, color="595959")

    allc = {s: by_split[s] for s in SPLITS}
    n = {s: len(allc[s]) for s in SPLITS}
    turns = sum(len(c["turns"]) for s in SPLITS for c in allc[s])
    facts = sum(len(c[k]) for s in SPLITS for c in allc[s] for k in ("assessment", "response", "education"))
    noise = sum(len(c["noise_events"]) for s in SPLITS for c in allc[s])
    flags = sum(len(c["risk_flags"]) for s in SPLITS for c in allc[s])
    errs = sum(counts[s][k] for s in SPLITS for k in counts[s])
    kpis = [("Training calls", n["train"], BLUE), ("Validation calls", n["val"], ORANGE), ("Transcript turns", turns, "5B9BD5"),
            ("Fact items", facts, "70AD47"), ("Logged noise events", noise, "FFC000"), ("Expected risk flags", flags, "C00000"),
            ("Validator findings", errs + sum(l.startswith("ERROR") for l in leaks), "548235")]
    for k, (name, val, color) in enumerate(kpis):
        col = 1 + k if k == 0 else 1 + k
        c1 = ws.cell(4, col, val)
        c1.font = Font(bold=True, size=22, color="FFFFFF")
        c1.fill = PatternFill("solid", fgColor=color)
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c2 = ws.cell(5, col, name)
        c2.font = Font(bold=True, color="FFFFFF", size=9)
        c2.fill = PatternFill("solid", fgColor=color)
        c2.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 38
    ws.row_dimensions[5].height = 26
    ws["A6"] = f"Validator: {sum(n.values())} calls checked, {errs} rule violations, {sum(l.startswith('ERROR') for l in leaks)} cross-split leaks. (Details: Validator sheet.)"
    ws["A6"].font = Font(italic=True, color="375623")
    for k in range(2, 8):
        ws.column_dimensions[get_column_letter(k)].width = 15
    b = Board(ws)

    def cnt(fn, split):
        return Counter(fn(c) for c in allc[split])

    # 1 split size
    b.table("1. Dataset size", ["Split", "Calls"], [["Training", n["train"]], ["Validation", n["val"]]], chart="pie", anchor_col=8, height=7.5, width=11)

    # 2 category vs target
    cat = {s: cnt(lambda c: c["setup"].get("scenario_category"), s) for s in SPLITS}
    b.table("2. Scenario categories (calls)", ["Category", "Training", "Validation", "Train target", "Val target"],
            [[k, cat["train"][k], cat["val"][k], CATEGORY_TARGET["train"][i], CATEGORY_TARGET["val"][i]] for i, k in enumerate(CATEGORY_ORDER)],
            series_cols=[2, 3], anchor_col=8, note="Training/Validation columns are the chart; target columns show the plan.")

    # 3 length
    ln = {s: cnt(lambda c: c["setup"].get("length_bucket"), s) for s in SPLITS}
    b.table("3. Call length", ["Bucket", "Training", "Validation", "Train %", "Val %"],
            [[k, ln["train"][k], ln["val"][k], round(ln["train"][k] / n["train"], 3), round(ln["val"][k] / n["val"], 3)] for k in ("short", "medium", "long")],
            series_cols=[2, 3], anchor_col=8, height=9, note="short >=65 lines, medium >=85, long >=105 (NA calls exempt).")
    for r in range(b.row - 10, b.row):
        for col in (4, 5):
            if isinstance(ws.cell(r, col).value, float):
                ws.cell(r, col).number_format = "0%"

    # 4 relationship
    rel = {s: cnt(lambda c: c["setup"].get("relationship_group") or "n/a", s) for s in SPLITS}
    keys = [k for k, _ in (rel["train"] + rel["val"]).most_common()]
    b.table("4. Who is calling", ["Caller", "Training", "Validation"], [[k, rel["train"][k], rel["val"][k]] for k in keys], anchor_col=8, height=10)

    # 5 age
    age = {s: cnt(lambda c: c["setup"].get("patient_age_group") or "n/a", s) for s in SPLITS}
    keys = [k for k, _ in (age["train"] + age["val"]).most_common()]
    b.table("5. Patient age group", ["Age", "Training", "Validation"], [[k, age["train"][k], age["val"][k]] for k in keys], anchor_col=8, height=9)

    # 6 agency type
    ag = {s: cnt(lambda c: c["setup"].get("agency_type"), s) for s in SPLITS}
    keys = [k for k, _ in (ag["train"] + ag["val"]).most_common()]
    held = [c for c in allc["val"] if c["setup"].get("agency_name") in ("Juniper Ridge Hospice", "Northstar Home Health")]
    b.table("6. Agency type", ["Agency type", "Training", "Validation"], [[k, ag["train"][k], ag["val"][k]] for k in keys], anchor_col=8, height=8,
            note=f"Validation includes {len(held)} calls from 2 held-out agencies never used in training (Juniper Ridge Hospice, Northstar Home Health).")

    # 7 risk flags
    fl = {s: Counter(f["risk_category"] for c in allc[s] for f in c["risk_flags"]) for s in SPLITS}
    keys = [k for k, _ in (fl["train"] + fl["val"]).most_common()]
    b.table("7. Risk flags expected", ["Risk category", "Training", "Validation"], [[k, fl["train"][k], fl["val"][k]] for k in keys], anchor_col=8, height=9,
            note="Flags are returned outside the summary text.")

    # 8 noise types
    nz = {s: Counter(e["type"] for c in allc[s] for e in c["noise_events"]) for s in SPLITS}
    keys = [k for k, _ in (nz["train"] + nz["val"]).most_common()]
    b.table("8. ASR / transcript noise events", ["Noise type", "Training", "Validation"], [[k, nz["train"][k], nz["val"][k]] for k in keys], anchor_col=8, height=9)

    # 9 noise level
    nl = {s: cnt(lambda c: c["setup"].get("asr_noise_level") or "n/a", s) for s in SPLITS}
    keys = [k for k in ("low", "medium", "high") if nl["train"][k] or nl["val"][k]] + [k for k in nl["train"] + nl["val"] if k not in ("low", "medium", "high")]
    b.table("9. ASR noise level per call", ["Level", "Training", "Validation"], [[k, nl["train"][k], nl["val"][k]] for k in dict.fromkeys(keys)], anchor_col=8, height=8)

    # 10 assessment fact types
    af = {s: Counter(i["item_type"] for c in allc[s] for i in c["assessment"]) for s in SPLITS}
    keys = [k for k, _ in (af["train"] + af["val"]).most_common()]
    b.table("10. Assessment fact types", ["Fact type", "Training", "Validation"], [[k, af["train"][k], af["val"][k]] for k in keys], anchor_col=8, height=9)

    # 11 medication status
    ms = {s: Counter({"ran": "ran out"}.get(i.get("medication_status"), i.get("medication_status")) for c in allc[s] for i in c["assessment"] if i["item_type"] == "medication") for s in SPLITS}
    keys = [k for k, _ in (ms["train"] + ms["val"]).most_common()]
    b.table("11. Medication status mentioned", ["Status", "Training", "Validation"], [[k or "unspecified", ms["train"][k], ms["val"][k]] for k in keys], anchor_col=8, height=9)

    # 12 actions & education
    ac = {s: Counter(i["action_type"] for c in allc[s] for i in c["response"]) for s in SPLITS}
    keys = [k for k, _ in (ac["train"] + ac["val"]).most_common()]
    b.table("12. Nurse response actions", ["Action", "Training", "Validation"], [[k, ac["train"][k], ac["val"][k]] for k in keys], anchor_col=8, height=10)
    ed = {s: Counter(i["education_type"] for c in allc[s] for i in c["education"]) for s in SPLITS}
    keys = [k for k, _ in (ed["train"] + ed["val"]).most_common()]
    b.table("13. Education given", ["Education type", "Training", "Validation"], [[k, ed["train"][k], ed["val"][k]] for k in keys], anchor_col=8, height=10)

    # 14 certainty
    ce = {s: Counter(i.get("certainty") for c in allc[s] for i in c["assessment"]) for s in SPLITS}
    b.table("14. Certainty of assessment facts", ["Certainty", "Training", "Validation"],
            [[k, ce["train"][k], ce["val"][k]] for k in ("stated", "unclear")], anchor_col=8, height=7.5, width=13,
            note="'unclear' marks hedged, conflicting or garbled facts that must not be stated as certain.")

    # 15 case checklist (the full width table + chart)
    rows = []
    for name, fn in CHECKLIST:
        a, v = sum(1 for c in allc["train"] if fn(c)), sum(1 for c in allc["val"] if fn(c))
        rows.append([name, a, v, "yes" if a and v else ("train only" if a else ("val only" if v else "MISSING"))])
    b.table("15. Case coverage checklist (calls containing the case)", ["Case", "Training", "Validation", "In both splits?"], rows,
            series_cols=[2, 3], anchor_col=8, height=18, width=22,
            note="Counts are calls, one call can cover several cases. Detected from the fact record, tags and noise log.")
    for r in range(b.row - 40, b.row):
        v = ws.cell(r, 4).value
        if v in ("yes",):
            ws.cell(r, 4).font = Font(color="548235", bold=True)
        elif v in ("MISSING", "val only", "train only"):
            ws.cell(r, 4).font = Font(color="C00000", bold=True)

    # 16 scenario subtypes (top 15 across both)
    st = {s: cnt(lambda c: c["setup"].get("subtype"), s) for s in SPLITS}
    top = [k for k, _ in (st["train"] + st["val"]).most_common(15)]
    b.table("16. Most common scenario subtypes", ["Subtype", "Training", "Validation"], [[k, st["train"][k], st["val"][k]] for k in top], anchor_col=8, height=11)

    # 17 transcript size
    def stat(split, fn):
        v = sorted(fn(c) for c in allc[split])
        return v[len(v) // 2], v[int(len(v) * 0.9)], v[-1]

    rows = []
    for label, fn in (("Turns per call", lambda c: len(c["turns"])), ("Characters per call", lambda c: len(c["transcript"]))):
        for s in SPLITS:
            med, p90, mx = stat(s, fn)
            rows.append([f"{label} ({LABEL[s]})", med, p90, mx])
    b.table("17. Transcript size (median / p90 / max)", ["Measure", "Median", "p90", "Max"], rows, chart=None)

    ws.freeze_panes = "A7"
    return ws


def main() -> int:
    by_split = {s: load_split(s) for s in SPLITS}
    wb = Workbook()
    wb.remove(wb.active)
    counts, calls_with, leaks = sheet_validator(wb, by_split)
    sheet_calls(wb, by_split)
    sheet_facts(wb, by_split)
    sheet_transcripts(wb, by_split)
    sheet_noise(wb, by_split)
    sheet_summaries(wb, by_split)
    sheet_analysis(wb, by_split, counts, calls_with, leaks)
    order = ["Analysis", "Calls", "Facts", "Transcripts", "Noise", "Summaries", "Validator"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}: " + ", ".join(f"{LABEL[s]} {len(by_split[s])}" for s in SPLITS))
    return 0


if __name__ == "__main__":
    sys.exit(main())
